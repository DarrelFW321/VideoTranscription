"""Google Translate's document UI; no API key or unofficial API endpoint."""
import logging
import json
import re
import time
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .common import digest, read_json, save_json
from .browser import browser_page

URL = "https://translate.google.com/?sl=ja&tl=en&op=docs&hl=en"


def export_document(segments, path):
    doc = Document()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ja-JP")
    lang.set(qn("w:eastAsia"), "ja-JP")
    doc.styles["Normal"].element.get_or_add_rPr().append(lang)
    for index, segment in enumerate(segments, 1):
        doc.add_paragraph(f"{index:08d} {segment['text']}")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    if path.stat().st_size > 10_000_000:
        raise ValueError("Document exceeds Google's 10 MB upload limit; split the transcript first")
    save_json(path.with_suffix(".mapping.json"), {"segments": segments, "fingerprint": digest(segments)})


def import_document(segments, path):
    values = {}
    doc = Document(path)
    for paragraph in doc.paragraphs:
        if not paragraph.text.strip():
            continue
        match = re.fullmatch(r"\s*([0-9]{8})\s+(.+?)\s*", paragraph.text, re.DOTALL)
        if not match:
            raise ValueError("Translated document changed a caption marker; review it manually")
        index, text = int(match[1]), match[2].strip()
        if index in values or not text or re.search(r"[\u3040-\u30ff]", text):
            raise ValueError(f"Duplicate, untranslated, or empty caption {index}")
        values[index] = text
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) != 2:
                raise ValueError("Translated document changed table layout; review it manually")
            number, text = row.cells[0].text.strip(), row.cells[1].text.strip()
            if not re.fullmatch(r"[0-9]{1,8}", number):
                raise ValueError(f"Translated document changed caption marker: {number!r}")
            index = int(number)
            if index in values or not text:
                raise ValueError("Duplicate marker or empty translated caption")
            if re.search(r"[\u3040-\u30ff]", text):
                raise ValueError(f"Caption {index} still contains Japanese kana; inspect the translation")
            values[index] = text
    if set(values) != set(range(1, len(segments) + 1)):
        raise ValueError("Translated document omitted/added caption markers; refusing misaligned subtitles")
    return [{**s, "japanese": s["text"],
             "text": "[inaudible]" if s["text"] == "[inaudible]" else values[i]}
            for i, s in enumerate(segments, 1)]


def translate_document(source, target, headed=False, timeout_ms=180000, profile=None, channel=None, stealth=False):
    from playwright.sync_api import sync_playwright
    target = Path(target)
    with sync_playwright() as p, browser_page(p, headed, profile, channel, stealth) as page:
        try:
            page.goto(URL, wait_until="domcontentloaded", timeout=60000)
            # Handle the normal consent dialog when Google's region requires it.
            reject = page.get_by_role("button", name="Reject all", exact=True)
            if reject.count() and reject.is_visible():
                reject.click()
            page.screenshot(path=str(target.with_suffix(".before.png")))
            page.locator('input[type="file"][accept*=".docx"]').set_input_files(str(Path(source).resolve()))
            page.get_by_role("button", name="Translate", exact=True).click(timeout=30000)
            download_button = page.get_by_role("button", name="Download translation", exact=True)
            download_button.wait_for(state="visible", timeout=timeout_ms)
            page.screenshot(path=str(target.with_suffix(".ready.png")))
            target.with_suffix(".page.txt").write_text(page.locator("body").inner_text(), encoding="utf-8")
            with page.expect_download(timeout=timeout_ms) as event:
                download_button.click()
            download = event.value
            logging.info("Google download filename: %s", download.suggested_filename)
            if download.failure():
                raise RuntimeError("Google document download failed")
            download.save_as(str(target))
        except Exception:
            page.screenshot(path=str(target.with_suffix(".failure.png")), full_page=True)
            target.with_suffix(".failure.txt").write_text(page.locator("body").inner_text(), encoding="utf-8")
            logging.error("Google web translation stopped. If Google shows a challenge, complete the upload manually; automation does not bypass challenges. Source: %s", source)
            raise


def parse_numbered(text, expected):
    matches = list(re.finditer(r"(?<!\d)([0-9]{8})\s+", text))
    values = {}
    for i, match in enumerate(matches):
        number = int(match[1])
        end = matches[i+1].start() if i+1 < len(matches) else len(text)
        value = text[match.end():end].strip()
        if number in values or not value or re.search(r"[\u3040-\u30ff]", value):
            raise ValueError("Duplicate/empty marker or untranslated Japanese")
        values[number] = value
    if set(values) != set(expected):
        raise ValueError("Google changed caption markers; refusing misaligned translation")
    return values


def json_payload(segments, start, end, context):
    entries = lambda items: [{"id": s["id"], "ja": s["text"]} for s in items]
    return {"payload": {"glossary": {}, "context_before": entries(segments[max(0, start-context):start]),
                        "previous_english": [], "targets": entries(segments[start:end]),
                        "context_after": entries(segments[end:end+context])}, "response": "{}"}


def json_batches(segments, batch_size=6, context=3, max_chars=4500):
    if batch_size < 1 or context < 0:
        raise ValueError("Invalid JSON batch/context size")
    batches, start = [], 0
    while start < len(segments):
        end, neighbors = min(len(segments), start + batch_size), context
        while True:
            payload = json_payload(segments, start, end, neighbors)
            text = json.dumps(payload, ensure_ascii=False, indent=2)
            if len(text) <= max_chars:
                break
            if neighbors:
                neighbors -= 1
            elif end - start > 1:
                end -= 1
            else:
                raise ValueError("One caption exceeds the JSON browser limit; split it in the reviewed transcript")
        batches.append((start, end, text))
        start = end
    return batches


def parse_json_translation(text, expected):
    def unique_keys(pairs):
        result = {}
        for k, v in pairs:
            if k in result:
                raise ValueError("Duplicate JSON key")
            result[k] = v
        return result
    try:
        result = json.loads(text, object_pairs_hook=unique_keys)
        targets = result["payload"]["targets"]
        if not isinstance(targets, list):
            raise ValueError("targets must be a list")
        values = {}
        for target in targets:
            key, value = target["id"], target["ja"]
            if key in values or not isinstance(value, str) or not value.strip() or re.search(r"[\u3040-\u30ff]", value):
                raise ValueError("Duplicate ID, untranslated Japanese, or empty caption")
            values[key] = value.strip()
        if set(values) != set(expected):
            raise ValueError("Caption IDs changed; refusing misaligned translation")
        return values
    except (KeyError, TypeError) as exc:
        raise ValueError("Google changed the JSON structure") from exc


def translate_text(segments, folder, resume=False, headed=False, input_format="json", batch_size=6, context=3, profile=None, channel=None, stealth=False):
    from playwright.sync_api import sync_playwright
    groups, current, size = [], [], 0
    for i, s in enumerate(segments, 1):
        line = f"{i:08d} {s['text']}"
        if len(line) > 3500:
            raise ValueError("A single caption exceeds the browser batch limit; split it in the reviewed transcript")
        if current and size + len(line) + 1 > 3500:
            groups.append(current)
            current, size = [], 0
        current.append((i, line))
        size += len(line) + 1
    if current:
        groups.append(current)
    prepared = []
    if input_format == "json":
        for start, end, text in json_batches(segments, batch_size, context):
            indices = list(range(start+1, end+1))
            prepared.append((indices, indices, text))
    else:
        for group in groups:
            first, last = group[0][0], group[-1][0]
            indices = list(range(max(1, first-2), min(len(segments), last+2)+1))
            text = "\n".join(f"{i:08d} {segments[i-1]['text']}" for i in indices)
            if len(text) > 4500:
                indices = [i for i, _ in group]
                text = "\n".join(line for _, line in group)
            prepared.append(([i for i, _ in group], indices, text))
    def parse(raw, indices):
        if input_format == "json":
            mapped = parse_json_translation(raw, [segments[i-1]["id"] for i in indices])
            return {i: mapped[segments[i-1]["id"]] for i in indices}
        return parse_numbered(raw, indices)
    values = {}
    with sync_playwright() as p, browser_page(p, headed, profile, channel, stealth) as page:
        try:
            for number, (targets, indices, text) in enumerate(prepared):
                path = folder / f"text_{number:05}_{digest(text)[:12]}.json"
                path.with_suffix(".input.txt").write_text(text, encoding="utf-8")
                if resume and path.exists():
                    result = parse(read_json(path)["response"], indices)
                else:
                    logging.info("Google browser translation block %s/%s (%s)", number+1, len(prepared), input_format)
                    page.goto(URL.replace("op=docs", "op=translate"), wait_until="domcontentloaded", timeout=60000)
                    reject = page.get_by_role("button", name="Reject all", exact=True)
                    if reject.count() and reject.is_visible():
                        reject.click()
                    page.locator('textarea[aria-label="Source text"]').fill(text, timeout=30000)
                    deadline, previous, stable = time.monotonic() + 60, None, 0
                    while time.monotonic() < deadline:
                        raw = "".join(page.locator('span[lang="en"] span[jsname="W297wb"]').all_text_contents())
                        try:
                            result = parse(raw, indices)
                        except ValueError:
                            stable = 0
                        else:
                            stable = stable + 1 if raw == previous else 0
                            if stable >= 2:
                                save_json(path, {"source": text, "response": raw})
                                break
                        previous = raw
                        page.wait_for_timeout(750)
                    else:
                        raise RuntimeError("Google did not return a complete aligned translation in 60 seconds")
                values.update({i: result[i] for i in targets})
        except Exception:
            page.screenshot(path=str(folder / "browser_failure.png"), full_page=True)
            (folder / "browser_failure.txt").write_text(page.locator("body").inner_text(), encoding="utf-8")
            raise
    return [{**s, "japanese": s["text"], "text": "[inaudible]" if s["text"] == "[inaudible]" else values[i]}
            for i, s in enumerate(segments, 1)]


def translate_google(segments, folder, resume=False, headed=False, method="text", input_format="json", batch_size=6, context=3, profile=None, channel=None, stealth=False):
    folder = Path(folder) / ("google_" + digest({"schema": 3, "segments": segments,
                                               "profile": str(Path(profile).resolve()) if profile else None,
                                               "channel": channel, "format": input_format, "stealth": stealth})[:16])
    folder.mkdir(parents=True, exist_ok=True)
    source, target = folder / "japanese.docx", folder / "english.docx"
    export_document(segments, source)
    if method == "text":
        result = translate_text(segments, folder, resume, headed, input_format, batch_size, context, profile, channel, stealth)
        export_document(result, folder / "english_from_text.docx")
        save_json(folder / "translations.json", result)
        return result
    if not (resume and target.exists()):
        logging.info("Uploading %s captions to Google Translate's Documents page", len(segments))
        translate_document(source, target, headed=headed, profile=profile, channel=channel, stealth=stealth)
    result = import_document(segments, target)
    save_json(folder / "translations.json", result)
    return result
